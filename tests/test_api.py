import io
from pathlib import Path
from pypdf import PdfWriter
from pptx import Presentation

SAMPLE = Path(__file__).resolve().parents[1] / "samples" / "Computer-networks.txt"


def upload(client):
    r = client.post("/api/documents", files={"file": ("lecture.txt", SAMPLE.read_bytes(), "text/plain")})
    assert r.status_code == 201, r.text
    return r.json()["id"]


def test_auth_and_session(client):
    assert client.get("/api/documents").status_code == 401
    data = {"email": "TEST@example.com", "password": "test-password-123", "name": "Test"}
    r = client.post("/api/auth/register", json=data)
    assert r.status_code == 201
    assert "httponly" in r.headers["set-cookie"].lower()
    assert r.json()["email"] == "test@example.com"
    assert "password_hash" not in r.json()
    assert client.post("/api/auth/register", json=data).status_code == 409
    assert client.post("/api/auth/logout").status_code == 204
    assert client.get("/api/auth/me").status_code == 401
    assert client.post("/api/auth/login", json={**data, "password": "wrong-password"}).status_code == 401
    assert client.post("/api/auth/login", json=data).status_code == 200


def test_full_study_workflow(account):
    c = account
    doc_id = upload(c)
    d = c.get(f"/api/documents/{doc_id}").json()
    assert d["summary"] and d["cards"] and d["source"]
    assert "questions" not in d
    answer = c.post(f"/api/documents/{doc_id}/ask", json={"question": "TCP reliable delivery"}).json()
    assert answer["found"] and answer["sources"][0]["page"] == 1
    assert "TCP" in answer["answer"]
    assert not c.post(f"/api/documents/{doc_id}/ask", json={"question": "astronaut photosynthesis"}).json()["found"]
    q = c.post(f"/api/documents/{doc_id}/quizzes").json()
    assert q["questions"] and "correct" not in q["questions"][0]
    assert c.post(f"/api/quizzes/{q['id']}/submit", json={"answers": []}).status_code == 422
    result = c.post(f"/api/quizzes/{q['id']}/submit", json={"answers": [0] * len(q["questions"])})
    assert result.status_code == 200 and 0 <= result.json()["score"] <= 100
    assert c.post(f"/api/quizzes/{q['id']}/submit", json={"answers": [0] * len(q["questions"])}).status_code == 409
    assert c.post(f"/api/documents/{doc_id}/reviews", json={"card_index": 0, "rating": "known"}).status_code == 201
    assert c.post(f"/api/documents/{doc_id}/reviews", json={"card_index": 100, "rating": "known"}).status_code == 422
    p = c.get("/api/progress").json()
    assert (p["documents"], p["quizzes"], p["reviews"]) == (1, 1, 1)
    assert c.delete(f"/api/documents/{doc_id}").status_code == 204
    assert c.get("/api/progress").json()["quizzes"] == 0


def test_ownership(account):
    c = account
    doc_id = upload(c)
    quiz = c.post(f"/api/documents/{doc_id}/quizzes").json()
    c.post("/api/auth/logout")
    c.post("/api/auth/register", json={"email": "other@example.com", "password": "test-password-123", "name": "Other"})
    assert c.get("/api/documents").json() == []
    for method, path, body in [
        ("get", f"/api/documents/{doc_id}", None),
        ("delete", f"/api/documents/{doc_id}", None),
        ("post", f"/api/documents/{doc_id}/ask", {"question": "What is TCP?"}),
        ("post", f"/api/quizzes/{quiz['id']}/submit", {"answers": [0]}),
    ]:
        kwargs = {"json": body} if body else {}
        assert getattr(c, method)(path, **kwargs).status_code == 404


def test_upload_validation(account):
    for name, data in [
        ("bad.pdf", b"not a pdf"),
        ("bad.exe", b"x" * 100),
        ("tiny.txt", b"tiny"),
        ("invalid.txt", b"\xff" * 100),
    ]:
        assert account.post("/api/documents", files={"file": (name, data)}).status_code == 422
    assert (
        account.post("/api/documents", files={"file": ("large.txt", b"x" * (10 * 1024 * 1024 + 1))}).status_code == 413
    )
    pdf = PdfWriter()
    pdf.add_blank_page(width=300, height=300)
    buffer = io.BytesIO()
    pdf.write(buffer)
    r = account.post("/api/documents", files={"file": ("scanned.pdf", buffer.getvalue())})
    assert r.status_code == 422 and "OCR" in r.json()["detail"]


def test_pptx_and_samples(account):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Networks"
    slide.placeholders[1].text = SAMPLE.read_text()
    buffer = io.BytesIO()
    deck.save(buffer)
    assert account.post("/api/documents", files={"file": ("lecture.pptx", buffer.getvalue())}).status_code == 201
    first = account.post("/api/documents/samples").json()
    assert len(first) == 2
    assert account.post("/api/documents/samples").json() == first
    arabic = next(d for d in first if "مبادئ" in d["title"])
    assert account.get(f"/api/documents/{arabic['id']}").json()["cards"]
    assert account.post(f"/api/documents/{arabic['id']}/ask", json={"question": "المفتاح الأساسي"}).json()["found"]


def test_csrf_and_rate_limit(client):
    assert client.post("/api/auth/logout", headers={"origin": "https://evil.example"}).status_code == 403
    for _ in range(15):
        client.post("/api/auth/login", json={"email": "nobody@example.com", "password": "test-password-123"})
    assert (
        client.post(
            "/api/auth/login", json={"email": "nobody@example.com", "password": "test-password-123"}
        ).status_code
        == 429
    )


def test_frontend_and_health(client):
    assert client.get("/").status_code == 200
    assert client.get("/assets/app.js").status_code == 200
    assert client.get("/api/health").json()["status"] == "ok"
    assert "/api/documents/{document_id}/ask" in client.get("/openapi.json").json()["paths"]


def test_text_pdf_and_expired_session(account):
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
    from backend.app.database import SessionLocal
    from backend.app.models import LoginSession

    pdf = PdfWriter()
    page = pdf.add_blank_page(width=612, height=792)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
    )
    stream = DecodedStreamObject()
    stream.set_data(
        b"BT /F1 12 Tf 40 740 Td (TCP provides reliable delivery by acknowledging data and retransmitting missing segments. DNS translates domain names into numerical addresses.) Tj ET"
    )
    page[NameObject("/Contents")] = pdf._add_object(stream)
    buffer = io.BytesIO()
    pdf.write(buffer)
    response = account.post("/api/documents", files={"file": ("network.pdf", buffer.getvalue(), "application/pdf")})
    assert response.status_code == 201
    result = account.get("/api/documents/" + str(response.json()["id"])).json()
    assert "TCP" in result["source"][0]["text"]
    with SessionLocal() as db:
        db.query(LoginSession).update({"expires": 0})
        db.commit()
    assert account.get("/api/documents").status_code == 401
