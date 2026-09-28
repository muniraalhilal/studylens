"""Deterministic, offline extractive study engine. No external data transfer."""

import io
import math
import re
import zipfile
from collections import Counter
from pathlib import Path
from pypdf import PdfReader
from pptx import Presentation

STOP = set(
    "the a an of to in is are and or for with that this as by on from what how why does explain about في من على إلى عن ما هو هي و أن هذا هذه التي الذي".split()
)


def tokens(text):
    return [w for w in re.findall(r"[^\W\d_]{2,}", text.lower()) if w not in STOP]


def extract(filename, data):
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        pdf = PdfReader(io.BytesIO(data))
        if pdf.is_encrypted or len(pdf.pages) > 200:
            raise ValueError("Use an unencrypted PDF with at most 200 pages / استخدمي ملف PDF غير مشفر حتى 200 صفحة")
        pages = [(i + 1, p.extract_text() or "") for i, p in enumerate(pdf.pages)]
    elif ext == ".pptx":
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 50_000_000:
                raise ValueError("Expanded file too large / حجم الملف بعد فك الضغط كبير")
        deck = Presentation(io.BytesIO(data))
        if len(deck.slides) > 200:
            raise ValueError("Maximum 200 slides / الحد الأقصى 200 شريحة")
        pages = [
            (i + 1, "\n".join(s.text for s in slide.shapes if s.has_text_frame)) for i, slide in enumerate(deck.slides)
        ]
    elif ext in {".txt", ".md"}:
        pages = [(1, data.decode("utf-8-sig"))]
    else:
        raise ValueError("Supported: PDF, PPTX, TXT, MD / أنواع الملفات المدعومة")
    chunks = []
    for page, content in pages:
        content = re.sub(r"[ \t]+", " ", content).strip()
        # Keep sentence boundaries and overlap by one complete sentence.
        units = re.split(r"(?<=[.!?؟])\s+|\n+", content)
        buffer = []
        for unit in units:
            unit = unit.strip()
            if not unit:
                continue
            if buffer and len("\n".join(buffer)) + len(unit) > 1100:
                chunks.append({"id": len(chunks) + 1, "page": page, "text": "\n".join(buffer)})
                buffer = [buffer[-1]] if len(buffer) > 1 else []
            buffer.append(unit)
        if buffer:
            chunks.append({"id": len(chunks) + 1, "page": page, "text": "\n".join(buffer)})
    if sum(len(c["text"]) for c in chunks) < 80:
        raise ValueError(
            "Not enough selectable text. Scanned PDFs need OCR first / النص غير كافٍ، الملفات الممسوحة تحتاج OCR"
        )
    if sum(len(c["text"]) for c in chunks) > 500_000:
        raise ValueError("Text limit exceeded / تجاوز النص الحد المسموح")
    return chunks


def generate(chunks):
    sentences = []
    for c in chunks:
        for s in re.split(r"(?<=[.!?؟])\s+|\n+", c["text"]):
            if 35 <= len(s.strip()) <= 500 and s.strip() not in [x["text"] for x in sentences]:
                sentences.append({"text": s.strip(), "page": c["page"], "chunk_id": c["id"]})
    if not sentences:
        sentences = [{**c, "chunk_id": c["id"]} for c in chunks[:8]]
    freq = Counter(tokens(" ".join(c["text"] for c in chunks)))
    ranked = sorted(
        sentences, key=lambda s: sum(freq[t] for t in set(tokens(s["text"]))) / math.sqrt(len(s["text"])), reverse=True
    )[:8]
    cards, questions = [], []
    terms = [t for t, _ in freq.most_common(80) if len(t) > 3]
    used_terms = set()
    for s in ranked:
        words = [w for w in terms if re.search(r"(?<!\w)" + re.escape(w) + r"(?!\w)", s["text"], re.I)]
        if not words:
            continue
        term = next((w for w in words if w not in used_terms), words[0])
        used_terms.add(term)
        prompt = re.sub(r"(?<!\w)" + re.escape(term) + r"(?!\w)", "_____", s["text"], flags=re.I)
        cards.append({"front": prompt, "back": term, "source": s})
        distractors = [t for t in terms if t != term and t not in tokens(s["text"])][:3]
        if len(distractors) == 3:
            options = sorted([term] + distractors)
            questions.append({"prompt": prompt, "options": options, "correct": options.index(term), "source": s})
    return {"mode": "offline-extractive", "summary": ranked[:5], "cards": cards, "questions": questions[:5]}


def retrieve(question, chunks):
    query = set(tokens(question))
    corpus = [Counter(tokens(c["text"])) for c in chunks]
    scored = []
    for c, counts in zip(chunks, corpus):
        score = sum(
            (1 + math.log(counts[t])) * (1 + math.log((len(chunks) + 1) / (1 + sum(t in d for d in corpus))))
            for t in query
            if counts[t]
        )
        if score > 0:
            scored.append((score, c))
    hits = []
    for _, chunk in sorted(scored, key=lambda pair: pair[0], reverse=True)[:3]:
        sentences = re.split(r"(?<=[.!?؟])\s+|\n+", chunk["text"])
        relevant = sorted(sentences, key=lambda sentence: len(query & set(tokens(sentence))), reverse=True)
        excerpt = "\n".join(s for s in relevant[:3] if query & set(tokens(s)))
        hits.append({**chunk, "text": excerpt or chunk["text"]})
    return {
        "found": bool(hits),
        "mode": "offline-extractive",
        "answer": "\n\n".join(c["text"] for c in hits),
        "sources": hits,
    }
